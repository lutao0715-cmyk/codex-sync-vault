# -*- coding: utf-8 -*-
"""三表联合核对：大表(直播全口径) vs 曝光侧表(推荐场景) 的 CTR 口径差异 + 33人全量统计"""
import openpyxl, statistics
from collections import defaultdict

A = 'C:/Users/win/Downloads/二台跨天 2.xlsx'          # 大表：直播全口径
E = 'C:/Users/win/Downloads/重点主播的开播&转化数据(1).xlsx'  # 曝光侧：推荐场景
B = 'C:/Users/win/Downloads/报名成功主播_累计数据_0914v2.xlsx'  # 名单：33人

def num(v):
    if v is None: return 0.0
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).replace(',', '').replace('%', '').strip())
    except Exception: return 0.0

def dstr(v):
    return '' if v is None else str(v).strip().replace('/', '-')[:10]

# ---------- 名单 ----------
wb = openpyxl.load_workbook(B, data_only=True, read_only=True)
ws = wb['关联结果明细']; ws.reset_dimensions()
rows = list(ws.iter_rows(values_only=True))
roster = {}
for r in rows[1:]:
    if r and r[0] is not None:
        roster[str(r[0]).strip()] = str(r[1]).strip()
wb.close()
print(f'名单: {len(roster)} 人')

# ---------- 大表聚合（按创作者ID） ----------
wb = openpyxl.load_workbook(A, data_only=True, read_only=True)
ws = wb['表格0']; ws.reset_dimensions()
it = ws.iter_rows(values_only=True)
hdr = next(it)
C = {'date': 0, 'cid': 1, 'name': 2, 'on_h': 19, 'play_dur': 20,
     'exp_c': 28, 'exp_u': 29, 'clk_c': 30, 'clk_u': 31,
     'play_c': 32, 'play_u': 33, 'inter_u': 43,
     'follow': 108, 'dj': 111, 'fans': 105, 'avg_min': 106}
agg = defaultdict(lambda: {'days': set(), 'lids': set(), 'exp_c': 0.0, 'exp_u_sum': 0.0,
                           'clk_c': 0.0, 'clk_u_sum': 0.0, 'play_c': 0.0, 'play_u_sum': 0.0,
                           'dur': 0.0, 'inter_u_sum': 0.0, 'follow': 0.0, 'dj': 0.0,
                           'fans': [], 'w_avg': 0.0, 'dirty': 0, 'rows': 0,
                           'exp_c_day': defaultdict(float), 'clk_c_day': defaultdict(float)})
for r in it:
    cid = str(r[C['cid']]).strip() if r[C['cid']] is not None else ''
    if cid not in roster:
        continue
    a = agg[cid]
    a['days'].add(dstr(r[C['date']]))
    a['rows'] += 1
    ec, cc = num(r[C['exp_c']]), num(r[C['clk_c']])
    if ec > 0 and cc > ec:          # 物理不可能的脏行
        a['dirty'] += 1
        continue
    a['exp_c'] += ec; a['clk_c'] += cc
    a['exp_u_sum'] += num(r[C['exp_u']]); a['clk_u_sum'] += num(r[C['clk_u']])
    a['play_c'] += num(r[C['play_c']]); a['play_u_sum'] += num(r[C['play_u']])
    a['dur'] += num(r[C['play_dur']]); a['inter_u_sum'] += num(r[C['inter_u']])
    a['follow'] += num(r[C['follow']]); a['dj'] += num(r[C['dj']])
    a['fans'].append(num(r[C['fans']]))
    a['w_avg'] += num(r[C['avg_min']]) * num(r[C['play_u']])
wb.close()
print(f'大表命中: {len(agg)} 人 | 明细行 {sum(a["rows"] for a in agg.values())} | 脏行 {sum(a["dirty"] for a in agg.values())}')

live = {}
for cid, a in agg.items():
    e = a['exp_c']
    live[cid] = {
        'name': roster[cid], 'days': len(a['days']), 'rows': a['rows'], 'dirty': a['dirty'],
        'exp_c': e, 'clk_c': a['clk_c'],
        'ctr_c': a['clk_c'] / e * 100 if e else 0,
        'ctr_u': a['clk_u_sum'] / a['exp_u_sum'] * 100 if a['exp_u_sum'] else 0,
        'exp_u_sum': a['exp_u_sum'], 'clk_u_sum': a['clk_u_sum'],
        'play_c': a['play_c'], 'play_u': a['play_u_sum'],
        'clk2play': a['play_u_sum'] / a['clk_u_sum'] * 100 if a['clk_u_sum'] else 0,   # 点击→播放转化
        'dur_h': a['dur'], 'inter_rate': a['inter_u_sum'] / a['play_u_sum'] * 100 if a['play_u_sum'] else 0,
        'follow': a['follow'], 'dj': a['dj'],
        'fans_net': (max(a['fans']) - min(a['fans'])) if a['fans'] else 0,
        'avg_min': a['w_avg'] / a['play_u_sum'] if a['play_u_sum'] else a['dur'] * 60 / a['play_u_sum'] if a['play_u_sum'] else 0,
        'fans_per_1k': (max(a['fans']) - min(a['fans'])) / e * 1000 if e else 0,
        'per_play': a['dur'] * 60 / a['play_c'] if a['play_c'] else 0,
    }

# ---------- 曝光侧表聚合（推荐场景，用户级） ----------
KEYMAP = {'date': ('行为产生时间',), 'aid': ('作者ID',), 'exp_u': ('曝光用户数',), 'clk_u': ('点击用户数',),
          'play_u': ('播放用户数',), 'dur': ('点击视频总播放时长',), 'min_pe': ('曝光人均播放时长',)}
def resolve(hdr):
    norm = [''.join(ch for ch in str(h or '') if ch not in '（()） ') for h in hdr]
    out = {}
    for k, pats in KEYMAP.items():
        for i, s in enumerate(norm):
            if any(s.startswith(p) for p in pats):
                out[k] = i; break
    return out

wb = openpyxl.load_workbook(E, data_only=True, read_only=True)
exp_agg = defaultdict(lambda: {'exp_u': 0.0, 'clk_u': 0.0, 'play_u': 0.0, 'dur': 0.0, 'days': set()})
name_by_id = {}
for sn in ['二台达人算法场景每日曝光', '二台达人首页feed场景每日曝光']:
    ws = wb[sn]; ws.reset_dimensions()
    I = None
    for r in ws.iter_rows(values_only=True):
        if not r or all(x is None for x in r):
            I = None; continue
        if len(r) > 2 and str(r[1]).strip() == '作者ID':
            I = resolve(r); continue
        if I is None or r[1] is None:
            continue
        aid = str(r[I['aid']]).strip()
        name_by_id[aid] = str(r[2] or '').strip()
        if aid not in roster:
            continue
        eu = num(r[I['exp_u']])
        d = exp_agg[aid]
        d['days'].add(dstr(r[I['date']]))
        d['exp_u'] += eu; d['clk_u'] += num(r[I['clk_u']]); d['play_u'] += num(r[I['play_u']])
        d['dur'] += num(r[I['dur']]) if 'dur' in I else num(r[I['min_pe']]) * eu
wb.close()
print(f'曝光侧命中: {len(exp_agg)} 人')

# ---------- 口径对账 ----------
print()
print(f"{'主播':<20}{'大表CTR次数%':>10}{'大表CTR用户%':>10}{'推荐CTR%':>9}{'推荐曝光用户':>10}{'大表曝光用户和':>12}{'推荐占比%':>9}{'点击→播放%':>10}")
print('-' * 96)
merged = []
for cid, L in live.items():
    x = exp_agg.get(cid, {'exp_u': 0, 'clk_u': 0, 'play_u': 0, 'dur': 0, 'days': set()})
    rec_ctr = x['clk_u'] / x['exp_u'] * 100 if x['exp_u'] else 0
    share = x['exp_u'] / L['exp_u_sum'] * 100 if L['exp_u_sum'] else 0
    merged.append(dict(cid=cid, **L, rec_exp_u=x['exp_u'], rec_clk_u=x['clk_u'],
                       rec_play_u=x['play_u'], rec_dur=x['dur'], rec_days=len(x['days']),
                       rec_ctr=rec_ctr, rec_share=share))
    print(f"{L['name']:<20}{L['ctr_c']:>10.2f}{L['ctr_u']:>10.2f}{rec_ctr:>9.2f}"
          f"{int(x['exp_u']):>10,}{int(L['exp_u_sum']):>12,}{share:>9.1f}{L['clk2play']:>10.1f}")
missing_exp = [roster[c] for c in roster if c not in agg]
print()
print('大表查无:', missing_exp if missing_exp else '无，33人全命中')
